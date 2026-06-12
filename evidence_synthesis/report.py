"""
Report assembly: per-hit Excel sheet (dashboard + embedded figures), an index sheet
for batch runs, and a SUMMARY.md. Presents every evidence line and its key numbers;
makes no automated one-gene-vs-two call.
"""
from __future__ import annotations
import os
import shutil
import time
from datetime import datetime

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

_HDR = Font(bold=True, size=12)
_SUB = Font(bold=True, size=10, color="FFFFFF")
_SUBFILL = PatternFill("solid", fgColor="4472C4")
_OK = PatternFill("solid", fgColor="E2EFDA")
_SKIP = PatternFill("solid", fgColor="F2F2F2")
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

_LINE_TITLES = {
    "cross_species": "Cross-species fused-protein hits",
    "taxonomy": "Taxonomic breadth",
    "read_through": "RNA-seq read-through junction",
    "read_level": "Read-level uniqueness (bridging reads)",
    "coexpression": "Single-cell co-occurrence",
    "bulk_expression": "Bulk expression",
}


def _line_metrics(name, res):
    """Return a list of (key, value) headline metrics for one evidence line."""
    if res.get("status") != "ok":
        return [("status", f"skipped — {res.get('reason', '')}")]
    s = res.get("summary", {})
    if name == "cross_species":
        rows = [("organisms (best 200 alignments)", s.get("organism_count_capped")),
                ("organisms (all alignments)", s.get("organism_count_uncapped")),
                ("best hit", f"{s.get('best_hit')} ({s.get('best_hit_organism')}), "
                 f"{s.get('best_pident')}% identity"),
                ("Gene 1 coverage", s.get("best_g1_frac")),
                ("Gene 2 coverage", s.get("best_g2_frac"))]
        if s.get("focal_best_hit"):
            rows.append(("melanogaster best hit", f"{s.get('focal_symbol')} "
                         f"(g1={s.get('focal_g1_frac')}, g2={s.get('focal_g2_frac')})"))
        else:
            rows.append(("melanogaster best hit", "none"))
        return rows
    if name == "taxonomy":
        rk = s.get("rank", "order")
        return [(f"{rk}s with hits", s.get(f"{rk}s_with_hits")),
                ("DB universe (n species)", s.get("db_universe_n")),
                ("overall hit rate %", s.get("overall_hitrate_pct"))]
    if name == "read_through":
        return [("distinct junctions", s.get("n_junctions")),
                ("top junction (coords)", s.get("top_junction")),
                ("top junction reads / UMIs", f"{s.get('top_reads')} / {s.get('top_umis')}"),
                ("reads at top / total", f"{s.get('reads_at_top')} / {s.get('total_bridging_reads')}")]
    if name == "read_level":
        return [("bridging reads", s.get("n_bridging_reads")),
                ("distinct UMIs", s.get("n_distinct_umis")),
                ("distinct barcodes", s.get("n_distinct_barcodes")),
                ("reads inspected", s.get("scan_inspected"))]
    if name == "coexpression":
        if s.get("merged_in_new_annotation"):
            return [("atlas gene", f"{s.get('atlas_var_g1')} (both LOCs map here)"),
                    ("new annotation", "MERGED into one gene — co-occurrence N/A")]
        return [("atlas genes", f"{s.get('atlas_var_g1')} / {s.get('atlas_var_g2')}"),
                ("cells co-express (frac)", s.get("frac_cells_coexpr")),
                ("Jaccard co-express (frac)", s.get("jaccard_coexpr"))]
    if name == "bulk_expression":
        ss = s.get("sample_set", "all")
        return [("samples", f"{s.get('n_samples')} ({ss})"),
                ("mean CPM g1", s.get("mean_cpm_g1")),
                ("mean CPM g2", s.get("mean_cpm_g2")),
                (f"log1p(CPM) Pearson r ({ss})", s.get("log1p_pearson_r"))]
    return [(k, v) for k, v in s.items()]


def add_hit_sheet(wb, bundle):
    hit, results, figs = bundle["hit"], bundle["results"], bundle["figs"]
    ws = wb.create_sheet(title=hit.name[:31])
    ws["A1"] = hit.fused_product or hit.name
    ws["A1"].font = _HDR
    meta = [("fused_protein", hit.fused_id), ("genes", f"{hit.gene_1} + {hit.gene_2}"),
            ("products", f"{hit.product_1} + {hit.product_2}"),
            ("composite_score", hit.composite_score),
            ("locus", f"{hit.chrom}:{hit.locus_window[1]:,}-{hit.locus_window[2]:,} ({hit.strand})")]
    r = 2
    for k, v in meta:
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v)
        r += 1

    r += 1
    for name in ["cross_species", "taxonomy", "read_through", "read_level",
                 "coexpression", "bulk_expression"]:
        res = results.get(name, {"status": "skipped", "reason": "not run"})
        c = ws.cell(r, 1, _LINE_TITLES[name])
        c.font = _SUB
        c.fill = _SUBFILL
        ws.cell(r, 2, "").fill = _SUBFILL
        ok = res.get("status") == "ok"
        r += 1
        for k, v in _line_metrics(name, res):
            ws.cell(r, 1, k)
            ws.cell(r, 2, v)
            ws.cell(r, 1).fill = _OK if ok else _SKIP
            ws.cell(r, 2).fill = _OK if ok else _SKIP
            r += 1
        r += 1

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 52
    # center every populated cell (labels + values)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=2):
        for c in row:
            c.alignment = _CENTER

    # embed figures to the right of the dashboard
    img_col = "D"
    img_row = 2
    for key in ("overview", "zoom"):
        p = figs.get(key)
        if p and os.path.exists(p):
            img = XLImage(p)
            img.anchor = f"{img_col}{img_row}"
            ws.add_image(img)
            img_row += 34
    return ws


def build_workbook(bundles, out_xlsx, cfg=None):
    if os.path.exists(out_xlsx):  # backup before overwrite (user verifies formatting)
        shutil.copy2(out_xlsx, out_xlsx + f".bak_{time.strftime('%Y%m%d_%H%M%S')}")
    wb = Workbook()
    idx = wb.active
    idx.title = "Index"
    headers = ["hit", "genes", "score", "organisms (best 200)", "organisms (all)",
               "orders hit", "bridging reads", "distinct UMIs", "coexpr (frac)"]
    for j, h in enumerate(headers, 1):
        c = idx.cell(1, j, h)
        c.font = _SUB
        c.fill = _SUBFILL
    for i, b in enumerate(bundles, 2):
        hit, res = b["hit"], b["results"]
        cs = res["cross_species"].get("summary", {})
        tx = res["taxonomy"].get("summary", {})
        rl = res["read_level"].get("summary", {})
        co = res["coexpression"].get("summary", {})
        row = [hit.name, f"{hit.gene_1}+{hit.gene_2}", hit.composite_score,
               cs.get("organism_count_capped"), cs.get("organism_count_uncapped"),
               tx.get(f"{tx.get('rank','order')}s_with_hits"),
               rl.get("n_bridging_reads"), rl.get("n_distinct_umis"),
               co.get("frac_cells_coexpr")]
        for j, v in enumerate(row, 1):
            idx.cell(i, j, v)
        add_hit_sheet(wb, b)
    for j in range(1, len(headers) + 1):
        idx.column_dimensions[get_column_letter(j)].width = 18
    for r in idx.iter_rows(min_row=1, max_row=idx.max_row, min_col=1, max_col=len(headers)):
        for c in r:
            c.alignment = _CENTER
    write_methods_sheet(wb, cfg)
    wb.move_sheet("Methods", -(len(wb.sheetnames) - 2))   # place Methods right after Index
    wb.save(out_xlsx)
    return out_xlsx


def write_methods_sheet(wb, cfg):
    ws = wb.create_sheet("Methods")
    rows = [
        ("HEADER", "GenoMiss evidence synthesis — methods & definitions", ""),
        ("", "Report generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("", "Question addressed", "Whether two adjacent annotated genes are actually one "
            "gene split in two by the annotation."),
        ("", "Note", "No automated one-gene-vs-two verdict is made; each evidence line is "
            "reported for you to weigh."),
        ("SECTION", "Inputs", ""),
        ("", "Protein sequences", getattr(cfg, "proteome", "")),
        ("", "Annotation used to find candidates (old)", getattr(cfg, "run_gff", "")),
        ("", "New annotation (figures)", getattr(cfg, "new_gff", "")),
        ("", "Cross-species protein database", getattr(cfg, "diamond_db", "")),
        ("", "Single-cell RNA-seq alignments", getattr(cfg, "bam", "")),
        ("", "Single-cell atlas", getattr(cfg, "atlas_h5ad", "")),
        ("", "Bulk RNA-seq counts", getattr(cfg, "bulk_matrix", "")),
        ("SECTION", "Key parameters", ""),
        ("", "Cross-species search", "DIAMOND blastp, expectation value 1e-5. 'Best 200 "
            "alignments' keeps only the 200 highest-scoring matches; 'all alignments' keeps "
            "up to 25,000."),
        ("", "Match filters", "At least 50 percent identity and 50 percent of the query "
            "aligned; the alignment must cross the gene one / gene two boundary by at least "
            "10 amino acids."),
        ("", "Bridging reads", "Kept only if uniquely mapped (mapping quality at least 250 "
            "and a single reported location) with at least 8 aligned bases on each side of "
            "the splice; independent molecules counted by unique molecular identifier."),
        ("", "Strandedness", "The library is reverse-stranded, so a read's transcript strand "
            "is the opposite of its alignment strand (this is the source of the apparent "
            "strand flip in genome viewers)."),
        ("", "Bulk expression", "Head samples only; counts-per-million normalized; "
            "correlation computed on natural-log(1 + value)."),
        ("SECTION", "Column / row definitions", ""),
        ("SUB", "Cross-species fused-protein hits", "Do other insects have a single protein "
            "matching both genes joined end to end?"),
        ("", "organisms (best 200 alignments)", "Number of species with a matching protein, "
            "using only the 200 best alignments (undercounts distant species)."),
        ("", "organisms (all alignments)", "Same, using all alignments — the true breadth."),
        ("", "best hit", "The single best-matching protein from any species: accession, "
            "species, and percent identity."),
        ("", "Gene 1 / Gene 2 coverage", "For that best hit only, the fraction of the gene "
            "one part and the gene two part of the joined query that the match covered. Both "
            "near 1 means one protein spans both genes."),
        ("", "melanogaster best hit", "The best-matching fruit fly (Drosophila melanogaster) "
            "protein — the best-annotated insect — shown as its gene symbol with its gene one "
            "and gene two coverage."),
        ("SUB", "Taxonomic breadth", ""),
        ("", "orders with hits", "Number of insect orders (major groups) containing at least "
            "one matching species."),
        ("", "database species (denominator)", "Total species in the search database."),
        ("", "overall hit rate", "Fraction of database species with a match."),
        ("SUB", "RNA-seq read-through junction", ""),
        ("", "distinct junctions", "Number of different splice points (donor-acceptor "
            "coordinate pairs) among the bridging reads."),
        ("", "top junction", "Coordinates of the splice supported by the most reads."),
        ("", "top junction reads / UMIs", "Read and independent-molecule support for it."),
        ("", "reads at top / total", "How concentrated the reads are on the single best "
            "splice — concentration suggests real read-through, scatter suggests artifacts."),
        ("SUB", "Read-level uniqueness", ""),
        ("", "bridging reads", "Uniquely-mapped reads aligning with one part in gene one and "
            "another in gene two across a splice."),
        ("", "distinct UMIs", "Independent original molecules among those reads (PCR "
            "duplicates removed via the unique molecular identifier)."),
        ("", "distinct barcodes", "Number of cells contributing those reads."),
        ("", "reads inspected", "Total alignment records examined over the locus before "
            "filtering."),
        ("SUB", "Single-cell co-occurrence", ""),
        ("", "atlas gene(s)", "The gene identifier(s) in the single-cell atlas the genes "
            "map to."),
        ("", "cells co-express (fraction)", "Of all cells, the fraction expressing both "
            "genes."),
        ("", "Jaccard co-express (fraction)", "Of cells expressing either gene, the fraction "
            "expressing both."),
        ("", "merged in new annotation", "Both genes map to one gene in the new annotation, "
            "so co-occurrence is not applicable."),
        ("SUB", "Bulk expression", ""),
        ("", "mean CPM gene 1 / gene 2", "Average expression in counts-per-million across "
            "the head samples."),
        ("", "log1p(CPM) Pearson correlation", "How strongly the two genes' expression "
            "tracks together across samples (counts-per-million, natural-log(1+value))."),
    ]
    r = 1
    for kind, a, b in rows:
        ca, cb = ws.cell(r, 1, a), ws.cell(r, 2, b)
        if kind == "HEADER":
            ca.font = _HDR
        elif kind == "SECTION":
            ca.font = _SUB
            ca.fill = _SUBFILL
            cb.fill = _SUBFILL
        elif kind == "SUB":
            ca.font = Font(bold=True)
        else:
            ca.font = Font(bold=True)
        ca.alignment = Alignment(vertical="top", wrap_text=True)
        cb.alignment = Alignment(vertical="top", wrap_text=True)
        r += 1
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 90
    return ws


def write_summary_md(hit, results, figs, path):
    L = []
    L.append(f"# Evidence synthesis — {hit.fused_product or hit.name}\n")
    L.append(f"**Fused protein:** `{hit.fused_id}`  ")
    L.append(f"**Genes:** {hit.gene_1} ({hit.product_1}) + {hit.gene_2} ({hit.product_2})  ")
    L.append(f"**Locus:** {hit.chrom}:{hit.locus_window[1]:,}-{hit.locus_window[2]:,} "
             f"({hit.strand})  ")
    L.append(f"**GenoMiss composite score:** {hit.composite_score}\n")
    L.append("Evidence below is presented without an automated one-gene-vs-two verdict.\n")
    for name in ["cross_species", "taxonomy", "read_through", "read_level",
                 "coexpression", "bulk_expression"]:
        res = results.get(name, {"status": "skipped", "reason": "not run"})
        L.append(f"## {_LINE_TITLES[name]}")
        if res.get("status") != "ok":
            L.append(f"_skipped — {res.get('reason','')}_\n")
            continue
        for k, v in _line_metrics(name, res):
            L.append(f"- **{k}:** {v}")
        L.append("")
    if figs.get("overview"):
        L.append(f"## Figures")
        L.append(f"- Locus overview (dual annotation): `{os.path.basename(figs['overview'])}`")
        if figs.get("zoom"):
            L.append(f"- Read-through junction zoom: `{os.path.basename(figs['zoom'])}`")
        L.append("")
    with open(path, "w") as fo:
        fo.write("\n".join(L))
    return path
